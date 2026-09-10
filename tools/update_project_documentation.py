from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


SOURCE = Path("QueryLens_Project_Documentation_Revised.docx")
OUTPUT = Path("QueryLens_Project_Documentation_Final.docx")


def replace_paragraph(document, starts_with, replacement):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(starts_with):
            paragraph.text = replacement
            return
    raise ValueError(f"Paragraph not found: {starts_with}")


def set_cell(cell, value, centered=False):
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(9)


document = Document(SOURCE)

# Screen descriptions now reflect user-visible CRUD operations.
screen_table = document.tables[1]
screen_updates = {
    "Connections": "Create, view, edit, delete, and select SQLite database connections",
    "Recommendations": "Review optimization advice, update its state, and delete saved recommendations",
    "Comparison History": "Search, rename, inspect, delete, and rerun saved comparison sessions",
}
for row in screen_table.rows[1:]:
    name = row.cells[0].text.strip()
    if name in screen_updates:
        set_cell(row.cells[1], screen_updates[name])

# The title is the only new comparison-session attribute; relationships remain unchanged.
entity_table = document.tables[2]
for row in entity_table.rows[1:]:
    if row.cells[0].text.strip() == "comparison_sessions":
        set_cell(row.cells[2],
                 "title is NOT NULL with a default value; one session owns zero or more comparison candidates")

replace_paragraph(
    document,
    "The repository layer currently exposes",
    "The repository and service layers expose complete create, read, update, and delete operations for "
    "three important entities: database connections, recommendations, and comparison sessions. "
    "The JavaFX screens expose the same operations with validation and confirmation for destructive actions."
)
replace_paragraph(
    document,
    "Current compliance note: persistence is meaningful",
    "Compliance result: the complete CRUD requirement is satisfied by Database Connection, Recommendation, "
    "and Comparison Session. Recommendation update is represented by its validated status transition, while "
    "Comparison Session update is represented by renaming its editable title."
)

crud_table = document.tables[3]
crud_values = {
    "Database Connection": ["Yes", "Yes", "Yes", "Yes"],
    "Query History": ["Yes", "Yes", "Not applicable", "Not yet"],
    "Recommendation": ["Yes", "Yes", "Status", "Yes"],
    "Comparison Session": ["Yes", "Yes", "Title", "Yes, cascades"],
}
for row in crud_table.rows[1:]:
    entity = row.cells[0].text.strip()
    if entity in crud_values:
        for index, value in enumerate(crud_values[entity], start=1):
            set_cell(row.cells[index], value, centered=True)

replace_paragraph(
    document,
    "DatabaseInitializer automatically creates",
    "DatabaseInitializer automatically creates the QueryLens history database directory, opens SQLite with "
    "foreign keys enabled, executes src/main/resources/database/schema.sql, and applies the backward-compatible "
    "comparison-session title migration when an existing database is opened."
)
replace_paragraph(
    document,
    "Current compliance note: the initializer creates tables",
    "DemoDataSeeder satisfies the data-seeding requirement. At application startup it creates a separate "
    "querylens-demo.db database, executes demo-schema.sql and demo-data.sql, loads deterministic Sailors, Boats, "
    "and Reserves records, creates useful demonstration indexes, and registers the database as QueryLens Demo. "
    "The process is idempotent: repeated startup uses INSERT OR IGNORE and the unique database path, so neither "
    "sample rows nor the saved connection are duplicated."
)

replace_paragraph(
    document,
    "Business-logic and persistence tests were executed",
    "Business-logic and persistence tests were executed with Maven using mvn test on 10 September 2026. "
    "Result: 42 tests run, 0 failures, 0 errors, and 0 skipped (BUILD SUCCESS)."
)

test_table = document.tables[5]
new_tests = [
    ("TC-13", "CRUD", "Create, read, update, and delete core entities",
     "Changes persist; missing records are rejected", "PASS"),
    ("TC-14", "Seeder", "Run DemoDataSeeder twice",
     "Demo tables and rows exist without duplicate data or connections", "PASS"),
]
for values in new_tests:
    row = test_table.add_row()
    for index, value in enumerate(values):
        set_cell(row.cells[index], value, centered=index in (0, 4))

requirements = document.tables[6]
for row in requirements.rows[1:]:
    requirement = row.cells[0].text.strip()
    if requirement == "Automated business tests":
        set_cell(row.cells[1], "42 passing Maven tests, including CRUD and repeated seeding")
    elif requirement == "Full CRUD for three entities":
        set_cell(row.cells[1], "Connections, recommendations, and comparison sessions provide complete CRUD")
        set_cell(row.cells[2], "Met", centered=True)
    elif requirement == "Schema and data seeder":
        set_cell(row.cells[1], "Initializer, migration, and idempotent Sailors/Boats/Reserves demo seeder")
        set_cell(row.cells[2], "Met", centered=True)

document.save(OUTPUT)
print(OUTPUT.resolve())
